from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

OUTPUT = Path("docs/tutorial/暴雪战网启动配置管理器使用教程.docx")
NAVY = "17365D"
BLUE = "2F75B5"
RED = "C00000"
LIGHT_BLUE = "EAF3F8"
LIGHT_GRAY = "F3F5F7"
DARK_GRAY = "404040"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_border(cell, color: str = RED, size: str = "16") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 160, start: int = 180, bottom: int = 160, end: int = 180) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_run(paragraph, text: str, *, bold: bool = False, color: str | None = None, size: int | None = None) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)


def add_body_paragraph(document: Document, text: str = "", *, space_after: int = 6, first_line: bool = False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15
    if first_line:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    if text:
        add_run(paragraph, text, size=10.5)
    return paragraph


def add_bullet(document: Document, text: str, level: int = 0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.1
    add_run(paragraph, text, size=10.5)
    return paragraph


def add_number(document: Document, text: str):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.1
    add_run(paragraph, text, size=10.5)
    return paragraph


def add_callout(document: Document, title: str, text: str, fill: str = LIGHT_BLUE, color: str = NAVY):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.5)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=color, size="10")
    set_cell_margins(cell, 160, 220, 160, 220)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    add_run(paragraph, title, bold=True, color=color, size=11)
    paragraph2 = cell.add_paragraph()
    paragraph2.paragraph_format.space_after = Pt(0)
    add_run(paragraph2, text, size=10.5)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_screenshot_placeholder(document: Document, number: int, title: str, note: str):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.5)
    cell.height = Cm(3.4)
    set_cell_shading(cell, "FAFAFA")
    set_cell_border(cell, color=RED, size="18")
    set_cell_margins(cell, 220, 260, 220, 260)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    add_run(paragraph, f"截图占位符 {number}：{title}", bold=True, color=RED, size=12)
    paragraph2 = cell.add_paragraph()
    paragraph2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph2.paragraph_format.space_after = Pt(0)
    add_run(paragraph2, note, color=DARK_GRAY, size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(document: Document, text: str, level: int = 1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for style_name, size, color in (("Title", 24, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 13, BLUE)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    if "Caption" in styles:
        styles["Caption"].font.name = "Microsoft YaHei"
        styles["Caption"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        styles["Caption"].font.size = Pt(9)
        styles["Caption"].font.color.rgb = RGBColor.from_string(DARK_GRAY)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    paragraph.add_run(" 页")


def build_document() -> None:
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(40)
    title.paragraph_format.space_after = Pt(12)
    add_run(title, "暴雪战网启动配置管理器", bold=True, color=NAVY, size=25)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    add_run(subtitle, "安装、配置 Windows 用户与创建桌面快捷方式教程", color=BLUE, size=15)

    add_callout(document, "本教程目标", "完成安装后，配置 Battle.net Launcher.exe 的路径、设置 Windows 用户密码、初始化本地 Windows 用户，并创建可直接启动对应配置的桌面快捷方式。", fill=LIGHT_BLUE)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(meta, f"文档版本：1.0　　日期：{date.today().isoformat()}", color=DARK_GRAY, size=9.5)
    document.add_page_break()

    add_heading(document, "一、完整操作流程", 1)
    add_body_paragraph(document, "推荐按照以下顺序操作：")
    flow = document.add_table(rows=1, cols=3)
    flow.alignment = WD_TABLE_ALIGNMENT.CENTER
    flow.autofit = False
    widths = [Cm(2.0), Cm(7.0), Cm(7.0)]
    headers = ["步骤", "操作", "结果"]
    for index, cell in enumerate(flow.rows[0].cells):
        cell.width = widths[index]
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, color="FFFFFF", size="6")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, headers[index], bold=True, color="FFFFFF", size=10)
    set_repeat_table_header(flow.rows[0])
    rows = [
        ("1", "打开安装程序", "安装器显示窗口并完成安装"),
        ("2", "修改 Battle.net Launcher.exe 路径", "工作目录自动更新为程序所在目录"),
        ("3", "输入密码并保存用户配置", "配置保存，凭据写入 Windows Credential Manager"),
        ("4", "初始化 Windows 账号", "按配置创建本地 Windows 用户"),
        ("5", "创建桌面快捷方式", "桌面出现当前配置专用快捷方式"),
        ("6", "启动当前配置", "以指定 Windows 用户启动 Battle.net"),
    ]
    for row_data in rows:
        row = flow.add_row()
        for index, value in enumerate(row_data):
            cell = row.cells[index]
            cell.width = widths[index]
            set_cell_shading(cell, "FFFFFF" if len(flow.rows) % 2 else LIGHT_GRAY)
            set_cell_border(cell, color="D9E2F3", size="4")
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            add_run(p, value, size=9.5)
    document.add_paragraph()

    add_heading(document, "二、打开安装程序", 1)
    add_number(document, "双击最新的安装包：")
    code = document.add_paragraph()
    code.paragraph_format.left_indent = Cm(0.8)
    add_run(code, r"D:\code\wow_runner\installer\dist\WowRunner-Setup.exe", bold=True, color=NAVY, size=10.5)
    add_number(document, "安装窗口打开后，确认安装位置。默认位置通常是当前用户的本地程序目录。")
    add_number(document, "勾选需要的选项，点击“开始安装”。")
    add_screenshot_placeholder(document, 1, "安装器主界面", "请插入安装器窗口截图，并用红框标出“开始安装”按钮。")
    add_number(document, "安装完成后，若勾选了“安装完成后启动配置管理器”，程序会自动打开配置管理器。")
    add_screenshot_placeholder(document, 2, "安装完成后自动打开配置管理器", "请插入配置管理器首次打开的截图。")

    add_heading(document, "三、修改 Battle.net Launcher.exe 路径", 1)
    add_body_paragraph(document, "配置器左侧是用户配置列表，右侧是当前选中配置的详细内容。先选择左侧的“wow”配置。")
    add_body_paragraph(document, "在“程序路径”一行点击“浏览”，选择实际安装的 Battle.net Launcher.exe。示例路径：")
    path = document.add_paragraph()
    path.paragraph_format.left_indent = Cm(0.8)
    add_run(path, r"C:\Program Files (x86)\Battle.net\Battle.net Launcher.exe", color=NAVY, size=10.5)
    add_callout(document, "重要提示", "程序路径可以按照你电脑上的实际安装位置修改。如果 Battle.net 安装在其他盘符或目录，请选择你电脑中真实存在的 Battle.net Launcher.exe。", fill="FFF2CC", color="9C6500")
    add_body_paragraph(document, "选择程序路径后，工作目录会自动更新为该 EXE 所在的前级目录，例如：")
    workdir = document.add_paragraph()
    workdir.paragraph_format.left_indent = Cm(0.8)
    add_run(workdir, r"C:\Program Files (x86)\Battle.net", color=NAVY, size=10.5)
    add_screenshot_placeholder(document, 3, "修改程序路径", "请插入配置器截图，并用红框圈出“程序路径”输入框和“浏览”按钮。")

    add_heading(document, "四、输入密码并保存用户配置", 1)
    add_body_paragraph(document, "在“密码（可选）”输入框中输入目标 Windows 用户的密码。教程示例可以使用：")
    password = document.add_paragraph()
    password.paragraph_format.left_indent = Cm(0.8)
    add_run(password, "123456", bold=True, color=RED, size=12)
    add_body_paragraph(document, "正式使用时请改成你自己的安全密码，不要直接使用示例密码。")
    add_body_paragraph(document, "点击“保存用户配置”后，会同时保存：")
    add_bullet(document, "程序路径、工作目录和参数列表。")
    add_bullet(document, "Windows 用户名配置。配置器只需要输入裸用户名，例如 wow1、wow-alt 或 wow_test，底层会自动保存为 .\\wow1、.\\wow-alt 或 .\\wow_test。")
    add_bullet(document, "如果密码框有内容，会覆盖保存该 profile 的旧凭据。")
    add_bullet(document, "密码不会写入 appsettings.json，而是保存到当前 Windows 用户的 Credential Manager。")
    add_callout(document, "覆盖密码说明", "如果之前已经保存过凭据，再输入新密码并点击“保存用户配置”，旧凭据会被新密码覆盖。如果密码框留空，则不会清空原来的凭据。", fill="E2F0D9", color="548235")
    add_screenshot_placeholder(document, 4, "输入密码并保存用户配置", "请插入配置器截图，用红框分别标出密码输入框和“保存用户配置”按钮。")

    add_heading(document, "五、初始化 Windows 用户", 1)
    add_body_paragraph(document, "保存用户配置后，点击“初始化 Windows 账号”。程序会请求管理员权限，并根据当前 profile 创建本机 Windows 用户。")
    add_body_paragraph(document, "关于密码的处理：")
    add_bullet(document, "通常情况下，初始化过程会直接使用刚才保存到 Credential Manager 的密码，不需要再次输入。")
    add_bullet(document, "如果 UAC 中切换成了另一个管理员账号，该管理员身份可能无法读取原用户的 Credential Manager。此时程序会再次弹出密码窗口，请输入同一个 Windows 用户密码。")
    add_bullet(document, "初始化只支持本机 Windows 用户，不支持域账号。")
    add_screenshot_placeholder(document, 5, "初始化 Windows 账号按钮", "请插入配置器截图，用红框标出“初始化 Windows 账号”按钮。")
    add_screenshot_placeholder(document, 6, "管理员权限确认或密码输入窗口", "请插入 UAC 或密码输入窗口截图。若系统没有再次询问密码，可将此处替换为“使用已保存凭据完成初始化”的截图。")
    add_callout(document, "本步骤的明确结论", "优先使用你在配置器中保存的密码；只有在 UAC 切换到其他管理员身份、无法读取原凭据时，才需要再次输入密码。", fill="E2F0D9", color="548235")

    add_heading(document, "六、创建桌面快捷方式", 1)
    add_body_paragraph(document, "确认 profile 已保存后，点击“创建桌面快捷方式”。")
    add_bullet(document, "如果程序路径是 Battle.net Launcher.exe，快捷方式名称为“暴雪战网 - 配置文件名.lnk”。例如配置名为 wow-alt 时，快捷方式为“暴雪战网 - wow-alt.lnk”。")
    add_bullet(document, "快捷方式会使用 Battle.net Launcher.exe 自己的图标。")
    add_bullet(document, "快捷方式会记住当前 profile 对应的 Windows 用户、参数和路径。")
    add_bullet(document, "不同 profile 可以创建多个快捷方式。")
    add_screenshot_placeholder(document, 7, "创建桌面快捷方式按钮", "请插入配置器截图，用红框标出“创建桌面快捷方式”按钮。")
    add_screenshot_placeholder(document, 8, "桌面上的快捷方式", "请插入桌面截图，用红框标出新生成的“暴雪战网 - 用户配置名”快捷方式。")

    add_heading(document, "七、启动当前配置", 1)
    add_body_paragraph(document, "有两种启动方式：")
    add_bullet(document, "回到配置管理器，点击“启动当前配置”。")
    add_bullet(document, "双击桌面上对应 profile 的快捷方式。")
    add_body_paragraph(document, "如果同时运行了其他 Battle.net 或 Agent 进程，Battle.net 可能会复用已有会话。测试不同 Windows 用户时，建议先关闭已有的 Battle.net、Agent 和 WoW 进程。")
    add_screenshot_placeholder(document, 9, "启动当前配置", "请插入点击“启动当前配置”后的截图，或插入 Battle.net 启动成功截图。")

    add_heading(document, "八、常见问题", 1)
    faq = [
        ("安装器双击没有窗口", "请使用最新的 131 MB 左右版本。旧版本是无交互的自解压安装包，已经替换为有可见安装界面的 WinForms 安装器。"),
        ("启动配置管理器出现 CMD 窗口", "请重新安装最新版。主程序现在使用 WinExe 输出类型，正常从快捷方式启动不会额外打开控制台窗口。"),
        ("工作目录是否需要手动填写", "选择程序路径后会自动截取 EXE 所在目录；如果业务需要，也可以手动修改工作目录。"),
        ("删除 Windows 账号会不会删除用户目录", "选择同时删除 Windows 账号时，会在十秒确认后删除账号及其 Windows 用户配置目录。只删除左侧配置时不会删除 Windows 账号。"),
        ("快捷方式图标为什么不同", "Battle.net Launcher.exe 的 profile 快捷方式优先读取目标 EXE 的图标；如果目标路径不存在，则使用暴雪战网启动配置管理器图标。"),
    ]
    for question, answer in faq:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_run(p, f"{question}：", bold=True, color=NAVY, size=10.5)
        add_run(p, answer, size=10.5)

    add_heading(document, "九、操作完成检查清单", 1)
    checklist = [
        "安装器正常打开并完成安装。",
        "Battle.net Launcher.exe 路径已修改为本机实际路径。",
        "工作目录已自动更新为程序所在目录。",
        "已输入密码并点击“保存用户配置”。",
        "Windows 用户已成功初始化。",
        "桌面已生成对应 profile 的快捷方式。",
        "快捷方式名称和图标符合预期。",
        "双击快捷方式可以启动对应配置。",
    ]
    for item in checklist:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        add_run(paragraph, item, size=10.5)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
